"""Editable PWC and EY buyer entity masters."""
from __future__ import annotations

from io import BytesIO
from typing import Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from loguru import logger
from pydantic import Field
from sqlalchemy.orm import Session

from app.api.eee_taxi_rates import check_edit_password
from app.database import get_db
from app.models import User
from app.services.auth import require_permission
from app.services.eee_taxi_client_master import (
    ClientMasterRow, ClientMasterRows, client_master_response, get_client_master, save_client_master,
)
from app.services.eee_taxi_rates import get_edit_password_hash

router = APIRouter(prefix="/api/eee-taxi/clients", tags=["eee-taxi-clients"])
_eee_user = require_permission("eee_taxi")


class ClientMasterUpdateIn(ClientMasterRows):
    edit_password: str = Field(min_length=1, max_length=128)


@router.get("")
def read_clients(
    client_profile: Literal["pwc", "ey"] = Query("pwc"),
    _: User = Depends(_eee_user),
    db: Session = Depends(get_db),
) -> dict:
    data = client_master_response(db, client_profile)
    data["has_edit_password"] = bool(get_edit_password_hash(db))
    return data


@router.put("")
def update_clients(
    body: ClientMasterUpdateIn,
    client_profile: Literal["pwc", "ey"] = Query(...),
    user: User = Depends(_eee_user),
    db: Session = Depends(get_db),
) -> dict:
    check_edit_password(db, body.edit_password, user)
    before = get_client_master(db, client_profile)
    save_client_master(db, client_profile, body.rows, user.email)
    after = get_client_master(db, client_profile)
    logger.info("{} client master saved by {}: {} entities", client_profile.upper(), user.email, len(after))
    data = client_master_response(db, client_profile)
    data["has_edit_password"] = bool(get_edit_password_hash(db))
    data["change_counts"] = {
        "added": len(after.keys() - before.keys()),
        "removed": len(before.keys() - after.keys()),
        "changed": sum(1 for key in before.keys() & after.keys() if before[key] != after[key]),
    }
    return data


@router.post("/import")
async def import_pwc_clients(
    file: UploadFile = File(...),
    edit_password: str = Form(..., min_length=1, max_length=128),
    client_profile: Literal["pwc"] = Query(...),
    user: User = Depends(_eee_user),
    db: Session = Depends(get_db),
) -> dict:
    """Replace the PWC entity master from the workbook's GSTIN/address/name rows."""
    check_edit_password(db, edit_password, user)
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Upload an Excel .xlsx workbook.")
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Workbook must be 10 MB or smaller.")
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = [row[0] for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True)]
        workbook.close()
        cells = [str(value).strip() if value is not None else "" for value in values]
        cells = [value for value in cells if value and value.casefold() != "row labels"]
        if len(cells) % 3:
            raise ValueError("Expected each entity to have GSTIN, invoice address, and entity name rows.")
        rows = [ClientMasterRow(gstin=cells[i], address=cells[i + 1], entity_name=cells[i + 2])
                for i in range(0, len(cells), 3)]
        validated = ClientMasterRows(rows=rows)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read PWC entity workbook: {exc}") from exc
    before = get_client_master(db, client_profile)
    save_client_master(db, client_profile, validated.rows, user.email)
    logger.info("PWC client master replaced from workbook by {}: {} entities", user.email, len(validated.rows))
    data = client_master_response(db, client_profile)
    data["has_edit_password"] = bool(get_edit_password_hash(db))
    data["change_counts"] = {"added": len(set(data_row.gstin for data_row in validated.rows) - before.keys()),
                             "removed": len(before.keys() - set(data_row.gstin for data_row in validated.rows)),
                             "imported": len(validated.rows)}
    return data
